# -*- coding: utf-8 -*-
"""把抓取到的 CSV 汇总成一份多工作表的 Excel，便于交作业。"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "C:/Users/ddead/WorkBuddy/2026-10-02-16-03-48/"
STATS = ["PTS", "FG%", "FT%", "3PM", "REB", "AST", "STL", "BLK", "TO"]
POS = ["PG", "SG", "SF", "PF", "C"]
COLS = ["Position", "Team", "OverallRank"] + [x for s in STATS for x in (s, s + "_Rank")]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
TITLE_FONT = Font(name="微软雅黑", size=13, bold=True, color="1F3864")
NOTE_FONT = Font(name="微软雅黑", size=10, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load(path):
    with open(path, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def write_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws.append(COLS)
    style_header(ws, 1, len(COLS))
    for r in rows:
        ws.append([r.get(c) for c in COLS])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(COLS)):
        for cell in row:
            cell.font, cell.border = BODY_FONT, BORDER
            if isinstance(cell.value, float):
                cell.number_format = "0.0"
                cell.alignment = Alignment(horizontal="center")
            elif isinstance(cell.value, int):
                cell.alignment = Alignment(horizontal="center")
    widths = [9, 8, 11] + [8] * (len(COLS) - 3)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    cur = load(ROOT + "defense_vs_position_2025-26.csv")
    prev = load(ROOT + "defense_vs_position_2024-25.csv")

    wb = Workbook()
    wb.remove(wb.active)

    # ---- 说明页 ----
    ws = wb.create_sheet("字段说明")
    info = [
        ("NBA Defense vs Position 数据集", ""),
        ("数据来源", "https://hashtagbasketball.com/nba-defense-vs-position"),
        ("数据口径", "每支球队每个位置（PG/SG/SF/PF/C）每 48 分钟让对手得到的平均数据"),
        ("抓取时间", "2026-10-02"),
        ("样本规模", "30 支球队 × 5 个位置 = 150 条/赛季"),
        ("", ""),
        ("字段", "含义"),
        ("Position", "位置：PG 控卫 / SG 分卫 / SF 小前 / PF 大前 / C 中锋"),
        ("Team", "球队三字母缩写（如 BOS、LAL、OKC）"),
        ("OverallRank", "综合名次 1–150，1 为让对手数据最低（即防守该位置最好）"),
        ("PTS / 3PM / REB / AST / STL / BLK / TO", "对手场均得分 / 三分命中 / 篮板 / 助攻 / 抢断 / 盖帽 / 失误"),
        ("FG% / FT%", "对手投篮命中率 / 罚球命中率"),
        ("<统计量>_Rank", "该统计量在全部 150 个「球队×位置」组合中的排名，1 为让对手该项数据最低"),
        ("", ""),
        ("备注", "网站未公开 OverallRank 的权重算法，此处按原站数值原样保留。"),
        ("备注", "数据由 ASP.NET GridView 服务端渲染，非 JSON 接口，抓取脚本见 fetch_hb.py。"),
    ]
    for r in info:
        ws.append(list(r))
    ws["A1"].font = TITLE_FONT
    for row in ws.iter_rows(min_row=7, max_row=7, max_col=2):
        for c in row:
            c.fill, c.font = HEAD_FILL, HEAD_FONT
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        for c in row:
            if c.font.size != 13 and c.row != 7:
                c.font = NOTE_FONT
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 62
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="center")

    # ---- 明细页 ----
    write_sheet(wb, "2025-26全季", cur)
    write_sheet(wb, "2024-25全季", prev)

    # ---- 球队×位置矩阵 ----
    ws = wb.create_sheet("各队位置矩阵_2025-26")
    teams = sorted({r["Team"] for r in cur})
    ws.append(["Team"] + [p + "_综合名次" for p in POS] + ["五位置名次均值", "对手PTS均值"])
    style_header(ws, 1, 7)
    grid = []
    for t in teams:
        ranks, pts = [], []
        for p in POS:
            hit = [r for r in cur if r["Team"] == t and r["Position"] == p]
            ranks.append(int(hit[0]["OverallRank"]) if hit else None)
            pts.append(float(hit[0]["PTS"]) if hit else None)
        grid.append((t, ranks, sum(ranks) / len(ranks), sum(pts) / len(pts)))
    grid.sort(key=lambda x: x[2])
    for t, ranks, avg, avgpts in grid:
        ws.append([t] + ranks + [round(avg, 1), round(avgpts, 1)])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=7):
        for i, cell in enumerate(row):
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(horizontal="center")
            if i >= 5 and isinstance(cell.value, float):
                cell.number_format = "0.0"
    for i, w in enumerate([9] + [14] * 5 + [17, 13], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"

    out = ROOT + "NBA_Defense_vs_Position_数据集.xlsx"
    wb.save(out)
    print("SAVED", out)
    print("sheets:", wb.sheetnames)
    print("矩阵前 5（防守最好）:", [g[0] for g in grid[:5]])
    print("矩阵后 5（防守最差）:", [g[0] for g in grid[-5:]])


if __name__ == "__main__":
    main()
